"""Phase 0 forensic inspection: regenerates the evidence behind docs/.

Dumps every source to page-aware text, then recomputes every derived fact
asserted in docs/initial_rules.md and tests/evaluation/golden_cases.json
directly from the workbook.

Usage:
    python scripts/inspect_sources.py --source-dir "<pack dir>" [--out build/inspection]
    python scripts/inspect_sources.py --self-check     # assert the Phase 0 facts

Nothing here is application logic. The application must derive severity and
outcomes from the documents at runtime; the small tables below are Phase 0
findings recorded so the docs are reproducible and falsifiable.
"""

from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import os
import pathlib
import sys

import openpyxl
import pymupdf

SNAPSHOT = dt.datetime(2026, 8, 16, 11, 0)  # README sheet, Asia/Kolkata
WORKBOOK = "ParcelPilot_Assessment_Data.xlsx"

# Phase 0 derived severities (docs/initial_rules.md R4). Recorded as inspection
# findings only -- application code MUST classify from Support Policy v3 s2.
PHASE0_DERIVED_SEVERITY = {
    "TKT-501": "P1", "TKT-502": "P2", "TKT-503": "P3",
    "TKT-504": "P3", "TKT-505": "P1",
}

# 24x7 first-response targets in minutes. Business-hour targets are deliberately
# absent: the pack never defines a business calendar (ADR-007, gap G1).
TARGETS_24X7 = {
    "SRC-01 v3 CURRENT": 30,
    "SRC-02 v2 DEPRECATED": 60,
    "SRC-05 Northstar agreement": 15,
}


def sha256(path: pathlib.Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def ts(value):
    if value in (None, ""):
        return None
    return value if isinstance(value, dt.datetime) else dt.datetime.fromisoformat(str(value))


def load_workbook(source_dir: pathlib.Path):
    wb = openpyxl.load_workbook(source_dir / WORKBOOK, data_only=True)

    def sheet(name):
        rows = list(wb[name].iter_rows(values_only=True))
        header = rows[0]
        return [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]

    return wb, sheet


def dump(source_dir: pathlib.Path, out: pathlib.Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    for path in sorted(source_dir.iterdir()):
        if path.suffix == ".pdf":
            doc = pymupdf.open(path)
            body = "".join(
                f"\n===== PAGE {i} of {doc.page_count} =====\n{page.get_text('text')}"
                for i, page in enumerate(doc, 1)
            )
        elif path.suffix == ".xlsx":
            wb = openpyxl.load_workbook(path, data_only=True)
            parts = []
            for ws in wb.worksheets:
                parts.append(f"\n===== SHEET {ws.title} rows={ws.max_row} cols={ws.max_column} =====")
                for row in ws.iter_rows(values_only=True):
                    if any(c is not None for c in row):
                        parts.append(" | ".join("" if c is None else str(c) for c in row))
            body = "\n".join(parts)
        else:
            continue
        dest = out / (path.stem[:40] + ".txt")
        dest.write_text(f"# FILE {path.name}\n# SHA256 {sha256(path)}\n{body}", encoding="utf-8")
        print(f"  wrote {dest.name}  sha256={sha256(path)[:12]}")


def cancellation(order, account_id_with_waiver="ACCT-001"):
    """SOP v4 s1 + agreement waiver. Returns (default_outcome, applied_outcome)."""
    booked, requested = ts(order["booked_at"]), ts(order["cancellation_requested_at"])
    status = order["status"]
    if status == "DELIVERED":
        default = "cannot cancel (DELIVERED)"
    elif status == "PICKED_UP":
        default = "do not cancel -> return-to-origin"
    elif requested is None:
        default = "no cancellation requested"
    else:
        minutes = (requested - booked).total_seconds() / 60
        default = "INR 0 (within 30 min)" if minutes <= 30 else "INR 250 (after 30 min)"
    applied = default
    if (
        order["account_id"] == account_id_with_waiver
        and status == "BOOKED"
        and requested is not None
    ):
        applied = "INR 0 (Northstar agreement SRC-05 s2 waiver)"
    return default, applied


def service_credit(order):
    """SOP v4 s2 default vs LumenWorks SRC-06 s3. Returns dict or None."""
    if not order["carrier_fault"] or order["customer_fault"]:
        return None
    actual = ts(order["pickup_actual_at"])
    reference = actual or SNAPSHOT  # ADR-006: null pickup accrues from snapshot
    delay_h = (reference - ts(order["pickup_window_end"])).total_seconds() / 3600
    fee = float(order["shipment_fee_inr"])
    return {
        "delay_hours": delay_h,
        "reference": "pickup_actual_at" if actual else "SNAPSHOT (accruing)",
        "default_sop_inr": min(500.0, 0.10 * fee) if delay_h > 2 else 0.0,
        "lumenworks_clause_inr": 300.0 if delay_h > 4 else 0.0,
    }


def sla_deadlines(ticket, plan, severity):
    """24x7 targets only. Business-hour targets are not computable (gap G1)."""
    if severity != "P1" or plan != "Enterprise":
        return {}
    created = ts(ticket["created_at"])
    out = {}
    for source, minutes in TARGETS_24X7.items():
        if "Northstar" in source and ticket["account_id"] != "ACCT-001":
            continue
        due = created + dt.timedelta(minutes=minutes)
        out[source] = {"due": due, "breached_by_minutes": (SNAPSHOT - due).total_seconds() / 60}
    return out


def report(source_dir: pathlib.Path) -> dict:
    wb, sheet = load_workbook(source_dir)
    accounts = {a["account_id"]: a for a in sheet("accounts")}
    orders, tickets = sheet("orders"), sheet("tickets")

    print(f"\nSnapshot {SNAPSHOT} ({SNAPSHOT.strftime('%A')})  "
          f"accounts={len(accounts)} orders={len(orders)} tickets={len(tickets)}")

    print("\n-- cancellation (SOP v4 s1) --")
    for o in orders:
        default, applied = cancellation(o)
        print(f"  {o['order_id']}  {accounts[o['account_id']]['account_name']:<20}"
              f"{o['status']:<11} default={default:<30} applied={applied}")

    print("\n-- failed-pickup service credit (SOP v4 s2 / SRC-06 s3) --")
    credits = {}
    for o in orders:
        c = service_credit(o)
        if c:
            credits[o["order_id"]] = c
            print(f"  {o['order_id']}  delay={c['delay_hours']}h via {c['reference']}  "
                  f"default_sop=INR {c['default_sop_inr']}  lumenworks=INR {c['lumenworks_clause_inr']}")
    if not credits:
        print("  (none eligible)")

    print("\n-- first-response SLA, 24x7 targets only --")
    breaches = []
    for t in tickets:
        if t["status"] != "open":
            continue
        severity = PHASE0_DERIVED_SEVERITY[t["ticket_id"]]
        plan = accounts[t["account_id"]]["plan"]
        deadlines = sla_deadlines(t, plan, severity)
        if not deadlines:
            print(f"  {t['ticket_id']}  sev={severity} plan={plan:<10} "
                  f"-> business-hours target, NOT computable (gap G1)")
            continue
        for source, d in deadlines.items():
            over = d["breached_by_minutes"]
            verdict = f"BREACHED by {over:g}m" if over > 0 else f"within target ({-over:g}m left)"
            print(f"  {t['ticket_id']}  sev={severity} {source:<28} due={d['due']}  {verdict}")
        applied = "SRC-05 Northstar agreement" if t["account_id"] == "ACCT-001" else "SRC-01 v3 CURRENT"
        if deadlines[applied]["breached_by_minutes"] > 0:
            breaches.append(t["ticket_id"])

    print("\n-- data-quality anomalies --")
    for o in orders:
        actual, requested = ts(o["pickup_actual_at"]), ts(o["cancellation_requested_at"])
        if actual and requested and requested > actual:
            print(f"  {o['order_id']}: cancellation requested {requested} AFTER pickup {actual}")
        if o["status"] == "BOOKED" and not actual and ts(o["pickup_window_end"]) < SNAPSHOT:
            print(f"  {o['order_id']}: window ended {ts(o['pickup_window_end'])}, still BOOKED at snapshot")
    print("  accounts with no contract_file:",
          [a for a, v in accounts.items() if not v["contract_file"]])
    print("  orphan fields (referenced by no document): premium_support, last_customer_message_at")

    return {"credits": credits, "breaches": breaches, "orders": orders,
            "tickets": tickets, "accounts": accounts}


def self_check(source_dir: pathlib.Path) -> None:
    """Assert the Phase 0 facts the docs depend on. Fails loudly on drift."""
    wb, sheet = load_workbook(source_dir)
    accounts = {a["account_id"]: a for a in sheet("accounts")}
    orders = {o["order_id"]: o for o in sheet("orders")}
    tickets = {t["ticket_id"]: t for t in sheet("tickets")}

    snap = wb["README"]["B2"].value
    assert snap == "2026-08-16 11:00 Asia/Kolkata", snap
    assert SNAPSHOT.strftime("%A") == "Sunday"

    # R2: cancellation outcomes
    assert cancellation(orders["ORD-1001"]) == ("INR 250 (after 30 min)",
                                                "INR 0 (Northstar agreement SRC-05 s2 waiver)")
    assert cancellation(orders["ORD-1002"])[1] == "do not cancel -> return-to-origin"
    assert cancellation(orders["ORD-2001"])[1] == "INR 250 (after 30 min)"
    assert cancellation(orders["ORD-3001"])[1] == "INR 0 (within 30 min)"
    assert cancellation(orders["ORD-4001"])[1] == "cannot cancel (DELIVERED)"

    # R3: ORD-2002 is the only credit-eligible order; agreement changes the amount
    eligible = [oid for oid, o in orders.items() if service_credit(o)]
    assert eligible == ["ORD-2002"], eligible
    c = service_credit(orders["ORD-2002"])
    assert c["delay_hours"] == 4.5 and c["reference"].startswith("SNAPSHOT")
    assert c["default_sop_inr"] == 240.0 and c["lumenworks_clause_inr"] == 300.0
    assert max(c["default_sop_inr"], c["lumenworks_clause_inr"]) <= 1000  # no manager approval

    # R5: TKT-501 breaches only under the agreement; TKT-505 breaches under v3
    d501 = sla_deadlines(tickets["TKT-501"], "Enterprise", "P1")
    assert d501["SRC-05 Northstar agreement"]["breached_by_minutes"] == 15
    assert d501["SRC-01 v3 CURRENT"]["breached_by_minutes"] == 0      # exactly on the line
    assert d501["SRC-02 v2 DEPRECATED"]["breached_by_minutes"] == -30  # deprecated hides it
    d505 = sla_deadlines(tickets["TKT-505"], "Enterprise", "P1")
    assert d505["SRC-01 v3 CURRENT"]["breached_by_minutes"] == 120
    assert "SRC-05 Northstar agreement" not in d505  # ACCT-004 has no agreement

    # Orphan / gap facts the docs assert
    assert accounts["ACCT-001"]["premium_support"] is True
    assert not accounts["ACCT-004"]["contract_file"]
    assert all(o["customer_fault"] is False for o in orders.values())
    assert {t["ticket_id"] for t in tickets.values() if t["historical_resolution"]} == {"TKT-450", "TKT-451"}

    print("self-check: all Phase 0 facts hold")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--source-dir", default=os.environ.get("PARCELPILOT_SOURCE_DIR"))
    ap.add_argument("--out", default="build/inspection")
    ap.add_argument("--self-check", action="store_true")
    args = ap.parse_args()

    if not args.source_dir:
        print("error: pass --source-dir or set PARCELPILOT_SOURCE_DIR", file=sys.stderr)
        return 2
    source_dir = pathlib.Path(args.source_dir)
    if not (source_dir / WORKBOOK).exists():
        print(f"error: {WORKBOOK} not found in {source_dir}", file=sys.stderr)
        return 2

    if args.self_check:
        self_check(source_dir)
        return 0

    print(f"dumping {source_dir} -> {args.out}")
    dump(source_dir, pathlib.Path(args.out))
    report(source_dir)
    self_check(source_dir)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
