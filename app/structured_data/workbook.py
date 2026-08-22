"""Normalizes ParcelPilot_Assessment_Data.xlsx into typed domain objects.

Every workbook datetime is naive (docs/data_dictionary.md anomaly #6); this
module localizes them to the snapshot's own timezone (ADR-005) rather than
leaving that assumption implicit downstream.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, tzinfo
from pathlib import Path

import openpyxl
from pydantic import ValidationError

from app.errors import MalformedWorkbookError, MissingSourceFileError
from app.models.structured import Account, Order, Ticket
from app.time.clock import parse_snapshot


@dataclass(frozen=True)
class WorkbookData:
    snapshot: datetime
    currency: str
    accounts: list[Account]
    orders: list[Order]
    tickets: list[Ticket]


def load_workbook(path: Path) -> WorkbookData:
    if not path.exists():
        raise MissingSourceFileError(path.name, str(path.parent))

    wb = openpyxl.load_workbook(path, data_only=True)
    snapshot_raw, currency = _read_readme(wb)
    snapshot = parse_snapshot(snapshot_raw)
    tz = snapshot.tzinfo
    assert tz is not None  # parse_snapshot always returns tz-aware

    for name in ("accounts", "orders", "tickets"):
        if name not in wb.sheetnames:
            raise MalformedWorkbookError(f"missing sheet {name!r}")

    accounts = [_build_account(row) for row in _sheet_rows(wb["accounts"])]
    account_ids = {a.account_id for a in accounts}

    orders = [_build_order(row, tz) for row in _sheet_rows(wb["orders"])]
    tickets = [_build_ticket(row, tz) for row in _sheet_rows(wb["tickets"])]

    for order in orders:
        if order.account_id not in account_ids:
            raise MalformedWorkbookError(
                f"order {order.order_id} references unknown account {order.account_id}"
            )
    for ticket in tickets:
        if ticket.account_id not in account_ids:
            raise MalformedWorkbookError(
                f"ticket {ticket.ticket_id} references unknown account {ticket.account_id}"
            )

    return WorkbookData(
        snapshot=snapshot, currency=currency, accounts=accounts, orders=orders, tickets=tickets
    )


def _read_readme(wb: openpyxl.Workbook) -> tuple[str, str]:
    if "README" not in wb.sheetnames:
        raise MalformedWorkbookError("no README sheet")
    kv: dict[str, str] = {}
    for row in wb["README"].iter_rows(values_only=True):
        if row and row[0] is not None and len(row) > 1 and row[1] is not None:
            kv[str(row[0]).strip()] = str(row[1]).strip()
    if "Dataset snapshot" not in kv:
        raise MalformedWorkbookError("README missing 'Dataset snapshot' row")
    return kv["Dataset snapshot"], kv.get("Currency", "")


def _sheet_rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise MalformedWorkbookError(f"sheet {ws.title!r} is empty")
    header = rows[0]
    if any(h is None for h in header):
        raise MalformedWorkbookError(f"sheet {ws.title!r} has a blank header cell")
    return [dict(zip(header, r, strict=True)) for r in rows[1:] if any(c is not None for c in r)]


def _parse_dt_optional(value: object, tz: tzinfo, *, field: str) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        naive = value
    else:
        try:
            naive = datetime.fromisoformat(str(value))
        except ValueError as exc:
            raise MalformedWorkbookError(f"{field}: unparseable datetime {value!r}") from exc
    return naive if naive.tzinfo is not None else naive.replace(tzinfo=tz)


def _parse_dt(value: object, tz: tzinfo, *, field: str) -> datetime:
    parsed = _parse_dt_optional(value, tz, field=field)
    if parsed is None:
        raise MalformedWorkbookError(f"{field}: required datetime is missing")
    return parsed


def _build_account(row: dict) -> Account:
    try:
        return Account(
            account_id=row["account_id"],
            account_name=row["account_name"],
            plan=row["plan"],
            status=row["status"],
            csm=row["csm"],
            contract_file=row.get("contract_file") or None,
            premium_support=bool(row["premium_support"]),
            notes=row.get("notes") or "",
        )
    except (KeyError, ValidationError) as exc:
        raise MalformedWorkbookError(f"invalid accounts row {row}: {exc}") from exc


def _build_order(row: dict, tz: tzinfo) -> Order:
    try:
        return Order(
            order_id=row["order_id"],
            account_id=row["account_id"],
            carrier=row["carrier"],
            status=row["status"],
            booked_at=_parse_dt(row["booked_at"], tz, field="booked_at"),
            pickup_window_start=_parse_dt(
                row["pickup_window_start"], tz, field="pickup_window_start"
            ),
            pickup_window_end=_parse_dt(row["pickup_window_end"], tz, field="pickup_window_end"),
            pickup_actual_at=_parse_dt_optional(
                row.get("pickup_actual_at"), tz, field="pickup_actual_at"
            ),
            shipment_fee_inr=float(row["shipment_fee_inr"]),
            carrier_fault=bool(row["carrier_fault"]),
            customer_fault=bool(row["customer_fault"]),
            cancellation_requested_at=_parse_dt_optional(
                row.get("cancellation_requested_at"), tz, field="cancellation_requested_at"
            ),
            notes=row.get("notes") or "",
        )
    except (KeyError, ValidationError) as exc:
        raise MalformedWorkbookError(f"invalid orders row {row}: {exc}") from exc


def _build_ticket(row: dict, tz: tzinfo) -> Ticket:
    try:
        return Ticket(
            ticket_id=row["ticket_id"],
            account_id=row["account_id"],
            created_at=_parse_dt(row["created_at"], tz, field="created_at"),
            status=row["status"],
            subject=row["subject"],
            description=row["description"],
            channel=row["channel"],
            assigned_to=row["assigned_to"],
            last_customer_message_at=_parse_dt(
                row["last_customer_message_at"], tz, field="last_customer_message_at"
            ),
            historical_resolution=row.get("historical_resolution") or None,
        )
    except (KeyError, ValidationError) as exc:
        raise MalformedWorkbookError(f"invalid tickets row {row}: {exc}") from exc
